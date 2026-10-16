# core/exporter.py
import math
import os

import openpyxl


def _clean_header(value):
    return str(value).replace('\n', '').replace('\r', '').strip()


def _is_blank(val):
    if val is None:
        return True
    if isinstance(val, str):
        return val.strip() in ('', '#N/A', 'N/A', 'NA', 'nan', 'NaN', 'None', '<NA>')
    if isinstance(val, float):
        return math.isnan(val)
    return False


def _to_number(val):
    if _is_blank(val):
        return None
    try:
        return round(float(val), 2)
    except (TypeError, ValueError):
        return None


def export_to_excel(original_path, df, output_path):
    """把 df 中的安全库存计算结果回填到原 Excel 的“安全库存”列，另存为新文件。"""
    # 不覆盖原文件
    if os.path.normpath(os.path.abspath(original_path)) == os.path.normpath(os.path.abspath(output_path)):
        root, ext = os.path.splitext(output_path)
        output_path = f'{root}_计算结果{ext}'

    # 定位 df 中的“安全库存”列
    safety_col = None
    for col in df.columns:
        if _clean_header(col) == '安全库存':
            safety_col = col
            break
    if safety_col is None:
        for col in df.columns:
            if '安全库存' in _clean_header(col):
                safety_col = col
                break
    if safety_col is None:
        raise ValueError('DataFrame 中未找到“安全库存”列')

    wb = openpyxl.load_workbook(original_path)

    # 定位工作表（优先名称含“安全库存”）
    ws = None
    for name in wb.sheetnames:
        if '安全库存' in str(name):
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    # 第 1 行为表头，定位“安全库存”列
    safety_col_idx = None
    for cell in ws[1]:
        if cell.value is not None and _clean_header(cell.value) == '安全库存':
            safety_col_idx = cell.column
            break

    if safety_col_idx is None:
        raise ValueError('原 Excel 中未找到“安全库存”列')

    # 数据从第 2 行开始，与 df 行顺序一一对应
    for i, val in enumerate(df[safety_col].tolist()):
        ws.cell(row=2 + i, column=safety_col_idx).value = _to_number(val)

    wb.save(output_path)
    return output_path